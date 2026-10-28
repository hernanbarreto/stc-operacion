#!/usr/bin/env python3
"""Valida el esquemático contra un log de eventos del ATS (formato eventos_*.xlsx: eventtype, sourcetime, description).

Comprueba que cada identificador de vía referenciado en el log (CDV, secciones, agujas, señales, BFU,
estaciones/AVC TGMT, andenes de DBO) se resuelve a un elemento del esquemático. Las reglas de resolución
son las mismas que usa la web (src/components/esquematico/resolver.ts).

Uso:
  python3 validar_log.py --esquematico ../../public/data/esquematico_l1.json --eventos eventos.xlsx [--json reporte.json]
Sale con código 1 si queda algún identificador sin resolver.
"""
from __future__ import annotations

import argparse
import collections
import csv
import json
import re
import sys

AREAS = ['PAN', 'ZAR', 'GOM', 'BOU', 'BAL', 'MOC', 'SLA', 'CAN', 'MER', 'PIN', 'ISA', 'SAL',
         'BAD', 'CUA', 'INS', 'SEV', 'CHP', 'JNA', 'TCY', 'OBS']


class Resolver:
    def __init__(self, esq: dict):
        self.cdv = {c['id']: c for c in esq['cdvs']}
        for c in esq['cdvs']:
            for a in c.get('alias', []):
                self.cdv[a] = c
        self.sec = {s['id']: s for s in esq['secciones']}
        self.agujas = {a['id']: a for a in esq['agujas']}
        self.senales = {s['id']: s for s in esq['senales']}
        self.avc = {}
        for a in esq['avc']:
            self.avc[a['id']] = a
            for al in a.get('alias', []):
                self.avc[al] = a
        self.andenes = set()
        for p in esq['andenes']:
            for k in p.get('ids', []):
                self.andenes.add(k)
        self.tgmt = {}
        for e in esq['estaciones']:
            self.tgmt[e['tgmt_id']] = ('estacion', e['codigo'])
        for a in esq['avc']:
            if a.get('tgmt_id'):
                self.tgmt[a['tgmt_id']] = ('avc', a['id'])
        self.bfu = set()
        for p in esq['andenes']:
            for b in p['bfu']:
                self.bfu.add(b + p['estacion'])

    def track(self, ident: str):
        """CDV físico o sección (virtual / de aguja), con o sin prefijo de área."""
        cands = [ident]
        if len(ident) > 3 and ident[:3] in AREAS:
            cands.append(ident[3:])
        for c in cands:
            if c in self.sec:
                return ('seccion', c)
            if c in self.cdv:
                return ('cdv', self.cdv[c]['id'])
        return None

    def signal(self, ident: str):
        base = re.sub(r'LRTB$', '', ident)
        return ('senal', base) if base in self.senales else None

    def switch(self, ident: str):
        return ('aguja', ident) if ident in self.agujas else None

    def bfu_(self, ident: str):
        return ('bfu', ident) if ident in self.bfu else None

    def avc_(self, ident: str):
        return ('avc', self.avc[ident]['id']) if ident in self.avc else None

    def anden(self, ident: str):
        return ('anden', ident) if ident in self.andenes else None

    def station(self, num: int):
        return self.tgmt.get(num)


PATTERNS = [
    ('cdv', re.compile(r'circuito de vía ([A-Z0-9_]+)'), 'track'),
    ('cdv', re.compile(r'ocupación en ([A-Z0-9_]+) el cuál'), 'track'),
    ('cdv', re.compile(r'Ocupación intempestiva en ([A-Z0-9_]+)'), 'track'),
    ('cdv', re.compile(r'Tren numerado como \d+ en el ([A-Z0-9_]+)$'), 'track'),
    ('aguja', re.compile(r'aguja ([A-Z0-9_]+)'), 'switch'),
    ('aguja', re.compile(r'Aguja ([A-Z0-9_]+) \('), 'switch'),
    ('senal', re.compile(r'[Ss]eñal ([A-Z0-9_]+)'), 'signal'),
    ('bfu', re.compile(r'^(BFU[A-Z0-9]+) \('), 'bfu_'),
]


def iter_rows(path: str):
    if path.lower().endswith('.csv'):
        with open(path, newline='', encoding='utf-8') as f:
            r = csv.reader(f)
            next(r, None)
            for row in r:
                yield row
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        next(it, None)
        for row in it:
            yield row


def validar(esq: dict, eventos_path: str):
    """Recorre el log y devuelve (n_eventos, vistos, sin_resolver, ejemplo) por categoría.
    ejemplo[(cat, id)] = (sourcetime, description) de la primera aparición de cada id sin resolver."""
    R = Resolver(esq)
    seen = collections.defaultdict(collections.Counter)      # categoría -> id -> ocurrencias
    unresolved = collections.defaultdict(collections.Counter)
    ejemplo = {}
    cmd_re = re.compile(r'(?:comando|Comando) (\w+) ([A-Z]{3}) (\S+)(?: (\S+))?')
    n = 0

    def check(cat, ident, ok, row):
        seen[cat][ident] += 1
        if not ok:
            unresolved[cat][ident] += 1
            ejemplo.setdefault((cat, ident), (str(row[1]), str(row[2])))

    for row in iter_rows(eventos_path):
        n += 1
        et, desc = str(row[0] or ''), str(row[2] or '')
        for cat, pat, fn in PATTERNS:
            for m in pat.finditer(desc):
                ident = m.group(1)
                check(cat, ident, getattr(R, fn)(ident) is not None, row)
        if et.startswith('TGMT_'):
            m = re.search(r'estación (\d+)', desc)
            if m:
                num = int(m.group(1))
                check('tgmt_estacion', num, R.station(num) is not None, row)
        if et == 'COMMAND_OPERATOR':
            m = cmd_re.search(desc)
            if m:
                verb, area, a1, a2 = m.groups()
                for ident in (a1, a2):
                    if not ident:
                        continue
                    if re.match(r'^AVC\w+$', ident):
                        check('cmd_avc', ident, R.avc_(ident) is not None, row)
                    elif re.match(rf'^({"|".join(AREAS)})\d+$', ident):
                        check('cmd_anden', ident, R.anden(ident) is not None, row)
                    elif re.match(r'^VA\d(PAN|OBS)$', ident) or ident.startswith('TGMT_'):
                        continue
                    elif re.match(r'^(S|BS)[0-9A-Z]+$', ident):
                        check('cmd_senal', ident, R.signal(ident) is not None, row)
                    elif re.match(r'^A[0-9A-Z]+$', ident) and verb not in ('ActivarSaltoEstacion',):
                        check('cmd_aguja_o_seccion', ident,
                              R.switch(ident) is not None or R.track(ident) is not None, row)
                    elif re.match(r'^C[0-9Z]', ident):
                        check('cmd_seccion', ident, R.track(ident) is not None, row)
    return n, seen, unresolved, ejemplo


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--esquematico', required=True)
    ap.add_argument('--eventos', required=True)
    ap.add_argument('--json')
    args = ap.parse_args()
    esq = json.load(open(args.esquematico, encoding='utf-8'))
    n, seen, unresolved, ejemplo = validar(esq, args.eventos)
    report = dict(eventos=n, categorias={})
    ok = True
    for cat in sorted(seen):
        tot = len(seen[cat])
        bad = len(unresolved[cat])
        if bad:
            ok = False
        report['categorias'][cat] = dict(ids=tot, resueltos=tot - bad,
                                         sin_resolver=dict(unresolved[cat].most_common()))
        print(f'{cat:22s} ids={tot:4d} resueltos={tot - bad:4d} sin_resolver={bad}'
              + (f'  {dict(unresolved[cat].most_common(12))}' if bad else ''))
    for (cat, ident), (t, d) in sorted(ejemplo.items(), key=lambda kv: str(kv[0])):
        print(f'  sin resolver {cat}/{ident}: {t} {d}')
    if args.json:
        json.dump(report, open(args.json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
